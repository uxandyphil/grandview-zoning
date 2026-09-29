"""
New business filings from the Ohio Secretary of State.

The Secretary of State's site turns away automated downloads, so its new business filing
reports are downloaded by hand (ohiosos.gov > Businesses > Business Reports) and uploaded
to sos-uploads/. Each run keeps the Grandview filings from any uploaded report (business
name, filing type, date, business address; no agent or contact names), adds them to
docs/business-filings.json, and deletes the statewide file so it isn't kept in the repo.
Column names are guessed from each report's header and logged, like the county files.
Runs from scrape.py.
"""

import csv
import io
import json
import re
from pathlib import Path

from county_permits import find_col, parse_date, title

UPLOADS = Path("sos-uploads")
OUT = Path("docs/business-filings.json")
GRANDVIEW = re.compile(r"^grandview\s*(heights|hts\.?|hgts)?$", re.I)
AGENT = (r"agent", r"contact", r"incorporator", r"registrant", r"filer", r"mail", r"principal")


def rows_from(path):
    """Header list and row dicts from a .csv/.txt/.xlsx/.xls report."""
    ext = path.suffix.lower()
    if ext == ".xlsx":
        from openpyxl import load_workbook
        grid = list(load_workbook(path, read_only=True, data_only=True).active.iter_rows(values_only=True))
    elif ext == ".xls":
        import xlrd
        sh = xlrd.open_workbook(path).sheet_by_index(0)
        grid = [sh.row_values(i) for i in range(sh.nrows)]
    else:
        text = path.read_bytes().decode("utf-8-sig", errors="replace")
        first = text.split("\n", 1)[0]
        delim = "\t" if "\t" in first else "|" if first.count("|") > first.count(",") else ","
        grid = list(csv.reader(io.StringIO(text), delimiter=delim))
    # the header is the first row with several filled cells (reports sometimes start with a title line)
    start = next((i for i, r in enumerate(grid[:15]) if sum(1 for v in r if str(v or "").strip()) >= 3), 0)
    header = [str(h or "").strip() for h in grid[start]]
    return header, [dict(zip(header, r)) for r in grid[start + 1:] if any(str(v or "").strip() for v in r)]


def columns(header):
    """Which columns hold what, preferring the business's own address over an agent's."""
    biz = [h for h in header if not any(re.search(a, h, re.I) for a in AGENT)]
    pick = lambda *pats: find_col(biz, *pats) or find_col(header, *pats)
    return dict(
        name=find_col(biz, r"business.?name", r"entity.?name", r"company", r"^name$", r"name"),
        number=find_col(header, r"charter", r"entity.?(no|num|id)", r"document.?(no|num)", r"filing.?(no|num)", r"^id$"),
        type=find_col(header, r"filing.?type", r"entity.?type", r"business.?type", r"transaction", r"^type$", r"type"),
        date=find_col(header, r"effective", r"filing.?date", r"filed", r"^date$", r"date"),
        address=pick(r"address.?1", r"street", r"address"),
        city=pick(r"city"),
        zip=pick(r"zip", r"postal"),
        county=pick(r"county"))


def run(state, geocode, now):
    data = json.loads(OUT.read_text()) if OUT.exists() else {"filings": [], "files": []}
    files = [f for f in sorted(UPLOADS.glob("*")) if f.suffix.lower() in (".csv", ".txt", ".xlsx", ".xls")] if UPLOADS.exists() else []
    known = {f["number"] or (f["name"], f["date"]): f for f in data["filings"]}
    for path in files:
        try:
            header, rows = rows_from(path)
        except Exception as err:
            print(f"Business filings: couldn't read {path.name} ({err})")
            continue
        cols = columns(header)
        print(f"Business filings: {path.name}, {len(rows)} rows, columns {header}")
        print(f"  using: {cols}")
        if not cols["name"] or not cols["city"]:
            print("  no business name or city column; leaving the file for a fix")
            continue
        found = 0
        for r in rows:
            if not GRANDVIEW.match(title(r.get(cols["city"]))):
                continue
            f = {"name": title(r.get(cols["name"])), "number": title(r.get(cols["number"])) if cols["number"] else "",
                 "type": title(r.get(cols["type"])) if cols["type"] else "", "date": parse_date(r.get(cols["date"])) if cols["date"] else None,
                 "address": title(r.get(cols["address"])) if cols["address"] else ""}
            key = f["number"] or (f["name"], f["date"])
            old = known.get(key, {})
            known[key] = {**f, "lat": old.get("lat"), "lon": old.get("lon")}
            found += 1
        print(f"  {found} Grandview filings")
        data["files"].append({"name": path.name, "rows": len(rows), "grandview": found, "added": now, "columns": header})
        path.unlink()   # the statewide report isn't kept; only the Grandview rows above
    data["filings"] = sorted(known.values(), key=lambda f: f["date"] or "", reverse=True)
    mapped = 0
    for f in data["filings"]:
        if f["address"] and f.get("lat") is None and not f.get("nomap"):
            c = geocode(f["address"] + ", Grandview Heights, OH")
            if c:
                f["lat"], f["lon"] = c[0], c[1]
                mapped += 1
            else:
                f["nomap"] = True
    if files or mapped or not OUT.exists():
        data["updated"] = now
        OUT.write_text(json.dumps(data, indent=1))
    print(f"Business filings: {len(data['filings'])} Grandview filings, {mapped} newly mapped")
